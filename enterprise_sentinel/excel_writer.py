from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from enterprise_sentinel.config import OUTPUT_COLUMNS


class ExcelReportWriter:
    """负责把标准化结果写入 Excel，并按报表场景做基础格式化。"""

    HEADER_FILL = PatternFill(fill_type="solid", fgColor="5E7D34")
    HEADER_FONT = Font(color="FFFFFF", bold=True)
    BODY_ALIGNMENT = Alignment(vertical="top", wrap_text=True)
    HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center")
    COLUMN_WIDTHS = {
        "A": 8,
        "B": 20,
        "C": 16,
        "D": 24,
        "E": 16,
        "F": 22,
        "G": 16,
        "H": 58,
        "I": 18,
        "J": 18,
        "K": 12,
        "L": 12,
        "M": 20,
    }

    def write(self, output_path: Path, frame: pd.DataFrame) -> None:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            export_frame = frame.reindex(columns=OUTPUT_COLUMNS)
            export_frame.to_excel(writer, sheet_name="监测结果", index=False)
            self._format_sheet(writer.book["监测结果"], export_frame)

    def _format_sheet(self, sheet, frame: pd.DataFrame) -> None:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = True

        for cell in sheet[1]:
            cell.fill = self.HEADER_FILL
            cell.font = self.HEADER_FONT
            cell.alignment = self.HEADER_ALIGNMENT

        for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
            for cell in row:
                cell.alignment = self.BODY_ALIGNMENT

        total_rows = max(len(frame) + 1, 2)
        for column_index in range(1, len(OUTPUT_COLUMNS) + 1):
            column_letter = get_column_letter(column_index)
            sheet.column_dimensions[column_letter].width = self.COLUMN_WIDTHS.get(column_letter, 16)

        for row_index in range(2, total_rows + 1):
            sheet.row_dimensions[row_index].height = 42
        sheet.row_dimensions[1].height = 24

