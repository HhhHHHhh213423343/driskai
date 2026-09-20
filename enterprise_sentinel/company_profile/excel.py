from __future__ import annotations

from datetime import datetime
from io import BytesIO
from math import ceil
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from enterprise_sentinel.company_profile.contract import MODULE_SPECS


TITLE_FILL = PatternFill("solid", fgColor="FF009A5A")
SECTION_FILL = PatternFill("solid", fgColor="FF166A45")
HEADER_FILL = PatternFill("solid", fgColor="FFEAF6F0")
ALT_FILL = PatternFill("solid", fgColor="FFF6FBF8")
BODY_FONT = Font(name="微软雅黑", color="FF1F2A24", size=10)
META_FONT = Font(name="微软雅黑", color="FF166A45", bold=True, size=10)
SOURCE_FONT = Font(name="微软雅黑", color="FF66736D", size=9)
HEADER_FONT = Font(name="微软雅黑", color="FF166A45", bold=True, size=10)
THIN_BORDER = Border(
    left=Side(style="thin", color="FFD7E2DA"),
    right=Side(style="thin", color="FFD7E2DA"),
    top=Side(style="thin", color="FFD7E2DA"),
    bottom=Side(style="thin", color="FFD7E2DA"),
)

SAMPLE_FREEZE_PANES = {
    "企业速览": "A39",
    "监管处罚": "A4",
    "动态监测": "A7",
    "工商变更": "A7",
    "高管信息": "A7",
    "股东信息": "A10",
    "对外投资企业": "A7",
    "控股子公司": "A7",
}

SAMPLE_COLUMN_WIDTHS = {
    "企业速览": [8.7, 108.7],
    "监管处罚": [5.7, 11.5, 6.9, 10.0, 13.0, 13.0, 52.5, 28.7, 23.7],
    "动态监测": [11.2, 71.2, 16.9, 10.2, 13.0, 20.0],
    "工商变更": [6.5, 13.1, 18.7, 56.0, 56.0],
    "高管信息": [8.7, 35.0, 37.5],
    "股东信息": [6.0, 42.5, 18.1, 11.2, 12.5, 13.7, 21.2],
    "对外投资企业": [6.0, 37.5, 11.2, 13.1, 12.5, 18.7, 16.2, 10.0, 13.0, 27.5],
    "控股子公司": [6.0, 38.7, 12.5, 11.2, 18.7, 13.7, 12.5, 10.0, 22.5],
}


def _rows_from_section(section: dict[str, Any]) -> tuple[list[str], list[list[Any]]]:
    columns = [str(item) for item in section.get("columns") or []]
    raw_rows = section.get("rows") or []
    if raw_rows and isinstance(raw_rows[0], dict):
        if not columns:
            columns = list(dict.fromkeys(key for row in raw_rows for key in row))
        return columns, [[row.get(column, "") for column in columns] for row in raw_rows]
    return columns, [list(row) for row in raw_rows]


def _style_data_area(sheet, start_row: int, end_row: int, max_column: int) -> None:
    for row in sheet.iter_rows(
        min_row=start_row,
        max_row=end_row,
        min_col=1,
        max_col=max_column,
    ):
        for cell in row:
            cell.font = BODY_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=True)


def _estimated_row_height(values: list[Any], widths: list[float]) -> float:
    lines = 1
    for index, value in enumerate(values):
        text = str(value or "")
        width = widths[index] if index < len(widths) else 18
        visual_lines = sum(
            max(1, ceil(len(part) / max(8, width * 1.45)))
            for part in text.splitlines() or [""]
        )
        lines = max(lines, visual_lines)
    return min(90.0, max(24.0, 15.0 + 13.5 * (lines - 1)))


def build_company_profile_workbook(
    company_name: str,
    normalized_data: dict[str, Any],
    captured_at: datetime,
) -> tuple[str, bytes]:
    workbook = Workbook()
    workbook.remove(workbook.active)
    local_time = captured_at.astimezone(ZoneInfo("Asia/Shanghai"))

    for spec in MODULE_SPECS:
        module = dict((normalized_data.get("modules") or {}).get(spec.key) or {})
        sections = list(module.get("sections") or [])
        if not sections:
            sections = [{"title": spec.title, "columns": list(spec.empty_columns), "rows": []}]
        max_columns = max(
            [2]
            + [max(1, len(_rows_from_section(section)[0])) for section in sections]
        )
        sheet = workbook.create_sheet(spec.title)
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = SAMPLE_FREEZE_PANES[spec.title]
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_columns)
        title = sheet.cell(1, 1, spec.title)
        title.fill = TITLE_FILL
        title.font = Font(name="微软雅黑", color="FFFFFFFF", bold=True, size=16)
        title.alignment = Alignment(horizontal="left", vertical="center")
        sheet.row_dimensions[1].height = 28

        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_columns)
        meta = sheet.cell(
            2,
            1,
            f"企业：{company_name}    提取日期：{local_time:%Y-%m-%d}",
        )
        meta.fill = HEADER_FILL
        meta.font = META_FONT
        meta.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[2].height = 19.5

        sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=max_columns)
        source = sheet.cell(
            3,
            1,
            f"来源：企业预警通 {module.get('source_url') or ''}",
        )
        source.font = SOURCE_FONT
        source.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[3].height = 22.5

        cursor = 5
        for section_index, section in enumerate(sections):
            if section.get("kind") == "note":
                sheet.merge_cells(
                    start_row=cursor,
                    start_column=1,
                    end_row=cursor,
                    end_column=max_columns,
                )
                note = sheet.cell(cursor, 1, section.get("text") or section.get("title") or "")
                note.font = META_FONT if section.get("emphasis") else SOURCE_FONT
                note.fill = HEADER_FILL if section.get("emphasis") else PatternFill(fill_type=None)
                note.alignment = Alignment(vertical="center", wrap_text=True)
                note.border = THIN_BORDER if not section.get("emphasis") else Border()
                sheet.row_dimensions[cursor].height = 21 if section.get("emphasis") else 36
                cursor += 1
                if section_index < len(sections) - 1 and sections[section_index + 1].get("kind") != "note":
                    cursor += 1
                continue

            columns, rows = _rows_from_section(section)
            columns = columns or list(spec.empty_columns)
            width = max(1, len(columns))
            if width > max_columns:
                max_columns = width
            sheet.merge_cells(
                start_row=cursor,
                start_column=1,
                end_row=cursor,
                end_column=width,
            )
            section_cell = sheet.cell(cursor, 1, section.get("title") or spec.title)
            section_cell.fill = SECTION_FILL
            section_cell.font = Font(name="微软雅黑", color="FFFFFFFF", bold=True, size=11)
            section_cell.alignment = Alignment(vertical="center")
            sheet.row_dimensions[cursor].height = 19.5
            cursor += 1

            for column_index, column in enumerate(columns, 1):
                cell = sheet.cell(cursor, column_index, column)
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.border = THIN_BORDER
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            sheet.row_dimensions[cursor].height = 25.5
            cursor += 1

            data_start = cursor
            for row_index, row in enumerate(rows):
                values = list(row) + [""] * max(0, len(columns) - len(row))
                for column_index, value in enumerate(values[: len(columns)], 1):
                    cell = sheet.cell(cursor, column_index, value)
                    if row_index % 2:
                        cell.fill = ALT_FILL
                sheet.row_dimensions[cursor].height = _estimated_row_height(
                    values[: len(columns)],
                    SAMPLE_COLUMN_WIDTHS[spec.title],
                )
                cursor += 1
            if not rows:
                sheet.cell(cursor, 1, "页面未显示记录")
                sheet.row_dimensions[cursor].height = 24
                cursor += 1
            _style_data_area(sheet, data_start, cursor - 1, len(columns))
            if section_index < len(sections) - 1:
                cursor += 1

        baseline_widths = SAMPLE_COLUMN_WIDTHS[spec.title]
        for column_index in range(1, max_columns + 1):
            width = baseline_widths[column_index - 1] if column_index <= len(baseline_widths) else 18
            sheet.column_dimensions[get_column_letter(column_index)].width = width
        sheet.auto_filter.ref = None
        sheet.print_title_rows = "1:6"
        sheet.page_setup.orientation = "landscape"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0

    output = BytesIO()
    workbook.save(output)
    data = output.getvalue()
    validate_company_profile_workbook(data)
    filename = f"{company_name}-企业全景信息-{local_time:%Y%m%d}.xlsx"
    return filename, data


def validate_company_profile_workbook(data: bytes) -> None:
    workbook = load_workbook(BytesIO(data), data_only=False)
    expected = [item.title for item in MODULE_SPECS]
    if workbook.sheetnames != expected:
        raise ValueError(f"Excel Sheet 顺序错误：{workbook.sheetnames}")
    formula_errors = {
        "#REF!",
        "#DIV/0!",
        "#VALUE!",
        "#NAME?",
        "#N/A",
        "#NUM!",
        "#NULL!",
        "#SPILL!",
        "#CALC!",
    }
    for sheet in workbook.worksheets:
        if sheet.max_row < 7:
            raise ValueError(f"{sheet.title} 缺少固定标题和业务区块。")
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    raise ValueError(f"{sheet.title}!{cell.coordinate} 不允许包含公式。")
                if isinstance(cell.value, str) and cell.value in formula_errors:
                    raise ValueError(f"{sheet.title}!{cell.coordinate} 包含公式错误。")
    if getattr(workbook, "_external_links", None):
        raise ValueError("Excel 不允许包含外部工作簿链接。")
